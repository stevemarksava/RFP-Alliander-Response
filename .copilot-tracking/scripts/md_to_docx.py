"""Convert the Alliander hardware BRD markdown into a shareable .docx.

Usage: python md_to_docx.py <input.md> <output.docx>
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tc_pr.append(shd)


def parse_frontmatter(lines):
    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1)
        fm_lines = lines[1:end]
        fm = {}
        for line in fm_lines:
            if ":" in line:
                key, _, value = line.partition(":")
                fm[key.strip()] = value.strip()
        return fm, lines[end + 1:]
    return {}, lines


def build_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, text in enumerate(header):
        hdr_cells[i].text = text
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
        set_cell_shading(hdr_cells[i], "2F5496")
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    doc.add_paragraph("")


def main():
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    text = src.read_text(encoding="utf-8")
    lines = text.splitlines()
    fm, body = parse_frontmatter(lines)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = fm.get("title", src.stem)
    subtitle = fm.get("description", "")
    author = fm.get("author", "")
    date = fm.get("ms.date", "")

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title)
    run.bold = True
    run.font.size = Pt(24)

    if subtitle:
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub_p.add_run(subtitle)
        run.italic = True
        run.font.size = Pt(13)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(f"{author}  |  {date}").font.size = Pt(10)
    doc.add_page_break()

    i = 0
    n = len(body)
    while i < n:
        line = body[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not body[i].strip().startswith("```"):
                code_lines.append(body[i])
                i += 1
            i += 1  # skip closing fence
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            doc.add_paragraph("")
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < n and body[i].strip().startswith("|"):
                table_lines.append(body[i].strip())
                i += 1
            header_cells = [c.strip() for c in table_lines[0].strip("|").split("|")]
            data_rows = []
            for row_line in table_lines[2:]:
                cells = [c.strip() for c in row_line.strip("|").split("|")]
                data_rows.append(cells)
            build_table(doc, header_cells, data_rows)
            continue

        if re.match(r"^\d+\.\s", stripped):
            items = []
            while i < n and re.match(r"^\d+\.\s", body[i].strip()):
                items.append(re.sub(r"^\d+\.\s", "", body[i].strip()))
                i += 1
            for item in items:
                doc.add_paragraph(item, style="List Number")
            continue

        if stripped.startswith("* "):
            items = []
            while i < n and body[i].strip().startswith("* "):
                items.append(body[i].strip()[2:])
                i += 1
            for item in items:
                doc.add_paragraph(item, style="List Bullet")
            continue

        # plain paragraph, possibly wrapped over multiple lines
        para_lines = [stripped]
        i += 1
        while i < n and body[i].strip() and not body[i].strip().startswith(("#", "|", "* ")) and not re.match(r"^\d+\.\s", body[i].strip()):
            para_lines.append(body[i].strip())
            i += 1
        doc.add_paragraph(" ".join(para_lines))

    doc.save(dst)
    print(f"Saved: {dst}")


if __name__ == "__main__":
    main()
